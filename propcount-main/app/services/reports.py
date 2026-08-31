from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Iterable, List, Optional

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func
from sqlmodel import Session, select

from app.enums import InvoiceKind, TransactionType
from app.models import Invoice, Lease, Object, Transaction, Unit
from app.schemas.reports import ReportResponse, ReportRow, ReportTotals
from app.services.realestate import get_user_object, get_user_unit
from app.utils.period import parse_period


def d(value) -> Decimal:
    return Decimal(str(value or 0))


def active_lease_unit_ids(
    session: Session,
    unit_ids: Iterable[int],
    month_start: date,
    month_end: date,
) -> set[int]:
    ids = list(unit_ids)

    if not ids:
        return set()

    leases = session.exec(
        select(Lease).where(
            Lease.unit_id.in_(ids),
            Lease.end_date >= month_start,
        )
    ).all()

    result = set()

    for lease in leases:
        if lease.start_date and lease.start_date > month_end:
            continue

        result.add(lease.unit_id)

    return result


def invoice_sum_for_object(
    session: Session,
    user_id: int,
    object_id: int,
    period: str,
    kind: str,
) -> Decimal:
    value = session.exec(
        select(func.coalesce(func.sum(Invoice.amount), 0))
        .join(Unit, Invoice.unit_id == Unit.id)
        .where(
            Invoice.user_id == user_id,
            Unit.object_id == object_id,
            Invoice.period == period,
            Invoice.kind == kind,
        )
    ).first()

    return d(value)


def invoice_sum_for_unit(
    session: Session,
    user_id: int,
    unit_id: int,
    period: str,
    kind: str,
) -> Decimal:
    value = session.exec(
        select(func.coalesce(func.sum(Invoice.amount), 0)).where(
            Invoice.user_id == user_id,
            Invoice.unit_id == unit_id,
            Invoice.period == period,
            Invoice.kind == kind,
        )
    ).first()

    return d(value)


def transaction_income_for_object(
    session: Session,
    user_id: int,
    object_id: int,
    date_from: date,
    date_to: date,
) -> Decimal:
    value = session.exec(
        select(func.coalesce(func.sum(Transaction.amount), 0)).where(
            Transaction.user_id == user_id,
            Transaction.object_id == object_id,
            Transaction.type == TransactionType.income.value,
            Transaction.date >= date_from,
            Transaction.date <= date_to,
        )
    ).first()

    return d(value)


def transaction_expense_for_object(
    session: Session,
    user_id: int,
    object_id: int,
    date_from: date,
    date_to: date,
) -> Decimal:
    value = session.exec(
        select(func.coalesce(func.sum(Transaction.amount), 0)).where(
            Transaction.user_id == user_id,
            Transaction.object_id == object_id,
            Transaction.type == TransactionType.expense.value,
            Transaction.date >= date_from,
            Transaction.date <= date_to,
        )
    ).first()

    return d(value)


def build_object_row(
    session: Session,
    user_id: int,
    obj: Object,
    period: str,
    date_from: date,
    date_to: date,
) -> ReportRow:
    units = session.exec(select(Unit).where(Unit.object_id == obj.id)).all()

    unit_ids = [unit.id for unit in units]

    occupied_ids = active_lease_unit_ids(
        session,
        unit_ids,
        date_from,
        date_to,
    )

    occupied_area = sum(
        (unit.area for unit in units if unit.id in occupied_ids),
        Decimal("0"),
    )

    free_area = max(obj.total_area - occupied_area, Decimal("0"))

    occupancy_percent = 0.0
    if obj.total_area > 0:
        occupancy_percent = round(float((occupied_area / obj.total_area) * 100), 2)

    rent_income = invoice_sum_for_object(
        session,
        user_id,
        obj.id,
        period,
        InvoiceKind.rent.value,
    )

    utility_income = invoice_sum_for_object(
        session,
        user_id,
        obj.id,
        period,
        InvoiceKind.utility.value,
    )

    transaction_income = transaction_income_for_object(
        session,
        user_id,
        obj.id,
        date_from,
        date_to,
    )

    expenses = transaction_expense_for_object(
        session,
        user_id,
        obj.id,
        date_from,
        date_to,
    )

    profit = rent_income + utility_income + transaction_income - expenses

    return ReportRow(
        entity_type="object",
        id=obj.id,
        name=obj.address,
        object_id=obj.id,
        object_address=obj.address,
        total_area=obj.total_area,
        occupied_area=occupied_area,
        free_area=free_area,
        occupancy_percent=occupancy_percent,
        rent_income=rent_income,
        utility_income=utility_income,
        transaction_income=transaction_income,
        expenses=expenses,
        profit=profit,
    )


def build_unit_row(
    session: Session,
    user_id: int,
    unit: Unit,
    obj: Object,
    period: str,
    date_from: date,
    date_to: date,
) -> ReportRow:
    occupied_ids = active_lease_unit_ids(
        session,
        [unit.id],
        date_from,
        date_to,
    )

    occupied_area = unit.area if unit.id in occupied_ids else Decimal("0")
    free_area = Decimal("0") if unit.id in occupied_ids else unit.area
    occupancy_percent = 100.0 if unit.id in occupied_ids else 0.0

    rent_income = invoice_sum_for_unit(
        session,
        user_id,
        unit.id,
        period,
        InvoiceKind.rent.value,
    )

    utility_income = invoice_sum_for_unit(
        session,
        user_id,
        unit.id,
        period,
        InvoiceKind.utility.value,
    )

    object_transaction_income = transaction_income_for_object(
        session,
        user_id,
        obj.id,
        date_from,
        date_to,
    )

    object_expenses = transaction_expense_for_object(
        session,
        user_id,
        obj.id,
        date_from,
        date_to,
    )

    ratio = Decimal("0")

    if obj.total_area > 0:
        ratio = unit.area / obj.total_area

    transaction_income = object_transaction_income * ratio
    expenses = object_expenses * ratio

    profit = rent_income + utility_income + transaction_income - expenses

    return ReportRow(
        entity_type="unit",
        id=unit.id,
        name=unit.number,
        object_id=obj.id,
        object_address=obj.address,
        total_area=unit.area,
        occupied_area=occupied_area,
        free_area=free_area,
        occupancy_percent=occupancy_percent,
        rent_income=rent_income,
        utility_income=utility_income,
        transaction_income=transaction_income,
        expenses=expenses,
        profit=profit,
    )


def build_totals(rows: List[ReportRow]) -> ReportTotals:
    total_area = sum((row.total_area for row in rows), Decimal("0"))
    occupied_area = sum((row.occupied_area for row in rows), Decimal("0"))
    free_area = sum((row.free_area for row in rows), Decimal("0"))

    occupancy_percent = 0.0

    if total_area > 0:
        occupancy_percent = round(float((occupied_area / total_area) * 100), 2)

    rent_income = sum((row.rent_income for row in rows), Decimal("0"))
    utility_income = sum((row.utility_income for row in rows), Decimal("0"))
    transaction_income = sum((row.transaction_income for row in rows), Decimal("0"))
    expenses = sum((row.expenses for row in rows), Decimal("0"))
    profit = sum((row.profit for row in rows), Decimal("0"))

    return ReportTotals(
        total_area=total_area,
        occupied_area=occupied_area,
        free_area=free_area,
        occupancy_percent=occupancy_percent,
        rent_income=rent_income,
        utility_income=utility_income,
        transaction_income=transaction_income,
        expenses=expenses,
        profit=profit,
    )


def get_landlord_report(
    session: Session,
    user_id: int,
    period: str,
    object_id: Optional[int] = None,
    unit_id: Optional[int] = None,
) -> ReportResponse:
    date_from, date_to = parse_period(period)

    rows: List[ReportRow] = []

    if unit_id is not None:
        unit = get_user_unit(session, user_id, unit_id)
        obj = session.get(Object, unit.object_id)

        if not obj:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Object not found",
            )

        rows.append(
            build_unit_row(
                session,
                user_id,
                unit,
                obj,
                period,
                date_from,
                date_to,
            )
        )

    elif object_id is not None:
        obj = get_user_object(session, user_id, object_id)

        units = session.exec(
            select(Unit).where(Unit.object_id == obj.id).order_by(Unit.number)
        ).all()

        rows = [
            build_unit_row(
                session,
                user_id,
                unit,
                obj,
                period,
                date_from,
                date_to,
            )
            for unit in units
        ]

    else:
        objects = session.exec(
            select(Object)
            .where(Object.user_id == user_id)
            .order_by(Object.created_at.desc())
        ).all()

        rows = [
            build_object_row(
                session,
                user_id,
                obj,
                period,
                date_from,
                date_to,
            )
            for obj in objects
        ]

    return ReportResponse(
        period=period,
        rows=rows,
        totals=build_totals(rows),
    )


DEFAULT_EXPORT_FIELDS = [
    "entity_type",
    "id",
    "name",
    "object_address",
    "total_area",
    "occupied_area",
    "free_area",
    "occupancy_percent",
    "rent_income",
    "utility_income",
    "transaction_income",
    "expenses",
    "profit",
]


FIELD_TITLES = {
    "entity_type": "Тип",
    "id": "ID",
    "name": "Название",
    "object_id": "ID объекта",
    "object_address": "Объект",
    "total_area": "Площадь",
    "occupied_area": "Занято",
    "free_area": "Свободно",
    "occupancy_percent": "Занятость, %",
    "rent_income": "Аренда",
    "utility_income": "Коммунальные",
    "transaction_income": "Прочие доходы",
    "expenses": "Расходы",
    "profit": "Прибыль",
}


def normalize_export_fields(fields: Optional[str]) -> List[str]:
    if not fields:
        return DEFAULT_EXPORT_FIELDS

    requested = [field.strip() for field in fields.split(",") if field.strip()]

    allowed = set(FIELD_TITLES.keys())

    result = [field for field in requested if field in allowed]

    return result or DEFAULT_EXPORT_FIELDS


def report_to_xlsx_bytes(
    report: ReportResponse,
    fields: Optional[str] = None,
) -> BytesIO:
    selected_fields = normalize_export_fields(fields)

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    ws.append([FIELD_TITLES.get(field, field) for field in selected_fields])

    header_fill = PatternFill("solid", fgColor="1F4E78")

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in report.rows:
        data = row.model_dump()
        ws.append([data.get(field) for field in selected_fields])

    ws.append([])

    totals = report.totals.model_dump()
    total_row = []

    for field in selected_fields:
        if field == "name":
            total_row.append("Итого")
        elif field in totals:
            total_row.append(totals[field])
        else:
            total_row.append("")

    ws.append(total_row)

    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            value = str(cell.value) if cell.value is not None else ""
            max_length = max(max_length, len(value))

        ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output
