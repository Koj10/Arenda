from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any, Dict, List, Optional

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlmodel import Session, select

from app.enums import InvoiceStatus, Plan, Role
from app.models import (
    Invoice,
    Lease,
    Object,
    Tenant,
    TenantProfile,
    Unit,
    UtilityBill,
)
from app.schemas.tenant_panel import (
    TenantInvoiceOut,
    TenantInvoicesResponse,
    TenantReportResponse,
    TenantSpaceOut,
    TenantSpacesResponse,
    TenantSubscriptionResponse,
)
from app.services.realestate import get_lease_status
from app.services.users import get_or_create_subscription
from app.utils.period import parse_period

EMPTY_STATE_MESSAGE = "Арендодатель должен добавить вашу компанию по ИНН"

INVOICE_REPORT_FIELDS = [
    "id",
    "period",
    "kind",
    "amount",
    "due_date",
    "computed_status",
    "object_address",
    "unit_number",
]

SPACE_REPORT_FIELDS = [
    "object_address",
    "unit_number",
    "area",
    "rent_monthly",
    "start_date",
    "end_date",
    "status",
]

FIELD_TITLES = {
    "id": "ID",
    "period": "Период",
    "kind": "Тип счёта",
    "amount": "Сумма",
    "due_date": "Оплатить до",
    "computed_status": "Статус",
    "object_address": "Объект",
    "unit_number": "Помещение",
    "area": "Площадь",
    "rent_monthly": "Аренда в месяц",
    "start_date": "Начало договора",
    "end_date": "Окончание договора",
    "status": "Статус",
}

TOTAL_FIELD_MAP = {
    "invoices": {
        "amount": "total_amount",
    },
    "spaces": {
        "area": "total_area",
        "rent_monthly": "monthly_rent",
    },
}


def decimal_zero() -> Decimal:
    return Decimal("0")


def get_tenant_profile(session: Session, user_id: int) -> Optional[TenantProfile]:
    return session.exec(
        select(TenantProfile).where(TenantProfile.user_id == user_id)
    ).first()


def get_matched_tenant_ids(session: Session, profile: TenantProfile) -> List[int]:
    tenants = session.exec(select(Tenant).where(Tenant.inn == profile.inn)).all()

    return [tenant.id for tenant in tenants]


def computed_invoice_status(invoice: Invoice) -> str:
    if (
        invoice.status == InvoiceStatus.pending.value
        and invoice.due_date < date.today()
    ):
        return "overdue"

    return invoice.status


def get_invoice_place(
    session: Session, invoice: Invoice
) -> tuple[Optional[str], Optional[str]]:
    object_address = None
    unit_number = None

    if invoice.unit_id:
        unit = session.get(Unit, invoice.unit_id)

        if unit:
            unit_number = unit.number
            obj = session.get(Object, unit.object_id)

            if obj:
                object_address = obj.address

    elif invoice.source_bill_id:
        bill = session.get(UtilityBill, invoice.source_bill_id)

        if bill:
            obj = session.get(Object, bill.object_id)

            if obj:
                object_address = obj.address

    return object_address, unit_number


def get_tenant_spaces(session: Session, user_id: int) -> TenantSpacesResponse:
    profile = get_tenant_profile(session, user_id)

    if not profile:
        return TenantSpacesResponse(
            matched=False,
            message=EMPTY_STATE_MESSAGE,
            spaces=[],
        )

    tenant_ids = get_matched_tenant_ids(session, profile)

    if not tenant_ids:
        return TenantSpacesResponse(
            matched=False,
            message=EMPTY_STATE_MESSAGE,
            spaces=[],
        )

    leases = session.exec(
        select(Lease)
        .where(Lease.tenant_id.in_(tenant_ids))
        .order_by(Lease.end_date.desc())
    ).all()

    spaces: List[TenantSpaceOut] = []

    for lease in leases:
        unit = session.get(Unit, lease.unit_id)

        if not unit:
            continue

        obj = session.get(Object, unit.object_id)

        if not obj:
            continue

        spaces.append(
            TenantSpaceOut(
                lease_id=lease.id,
                status=get_lease_status(lease.end_date),
                start_date=lease.start_date,
                end_date=lease.end_date,
                rent_monthly=lease.rent_monthly,
                unit_id=unit.id,
                unit_number=unit.number,
                unit_area=unit.area,
                object_id=obj.id,
                object_address=obj.address,
            )
        )

    return TenantSpacesResponse(
        matched=True,
        message=None,
        spaces=spaces,
    )


def get_tenant_invoices(
    session: Session,
    user_id: int,
    status_filter: Optional[str] = None,
) -> TenantInvoicesResponse:
    profile = get_tenant_profile(session, user_id)

    if not profile:
        return TenantInvoicesResponse(
            matched=False,
            message=EMPTY_STATE_MESSAGE,
            invoices=[],
        )

    tenant_ids = get_matched_tenant_ids(session, profile)

    if not tenant_ids:
        return TenantInvoicesResponse(
            matched=False,
            message=EMPTY_STATE_MESSAGE,
            invoices=[],
        )

    statement = select(Invoice).where(Invoice.tenant_id.in_(tenant_ids))

    if status_filter:
        status_value = status_filter.lower()

        if status_value == "all":
            pass
        elif status_value == InvoiceStatus.pending.value:
            statement = statement.where(
                Invoice.status == InvoiceStatus.pending.value,
                Invoice.due_date >= date.today(),
            )
        elif status_value == InvoiceStatus.paid.value:
            statement = statement.where(Invoice.status == InvoiceStatus.paid.value)
        elif status_value == "overdue":
            statement = statement.where(
                Invoice.status == InvoiceStatus.pending.value,
                Invoice.due_date < date.today(),
            )
        else:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Invalid status filter",
            )

    statement = statement.order_by(Invoice.due_date.desc())

    invoices = session.exec(statement).all()

    invoices_out: List[TenantInvoiceOut] = []

    for invoice in invoices:
        object_address, unit_number = get_invoice_place(session, invoice)

        invoices_out.append(
            TenantInvoiceOut(
                id=invoice.id,
                kind=invoice.kind,
                period=invoice.period,
                amount=invoice.amount,
                due_date=invoice.due_date,
                status=invoice.status,
                computed_status=computed_invoice_status(invoice),
                object_address=object_address,
                unit_number=unit_number,
            )
        )

    return TenantInvoicesResponse(
        matched=True,
        message=None,
        invoices=invoices_out,
    )


def empty_report_totals(tab: str) -> Dict[str, Any]:
    if tab == "invoices":
        return {
            "count": 0,
            "total_amount": decimal_zero(),
            "paid_amount": decimal_zero(),
            "pending_amount": decimal_zero(),
            "overdue_amount": decimal_zero(),
        }

    return {
        "count": 0,
        "total_area": decimal_zero(),
        "monthly_rent": decimal_zero(),
    }


def build_tenant_invoice_report_rows(
    session: Session,
    tenant_ids: List[int],
    period: str,
) -> tuple[List[Dict[str, Any]], Dict[str, Any]]:
    invoices = session.exec(
        select(Invoice)
        .where(
            Invoice.tenant_id.in_(tenant_ids),
            Invoice.period == period,
        )
        .order_by(Invoice.due_date.desc())
    ).all()

    rows: List[Dict[str, Any]] = []

    total_amount = decimal_zero()
    paid_amount = decimal_zero()
    pending_amount = decimal_zero()
    overdue_amount = decimal_zero()

    for invoice in invoices:
        object_address, unit_number = get_invoice_place(session, invoice)
        computed_status = computed_invoice_status(invoice)

        rows.append(
            {
                "id": invoice.id,
                "period": invoice.period,
                "kind": invoice.kind,
                "amount": invoice.amount,
                "due_date": invoice.due_date,
                "computed_status": computed_status,
                "object_address": object_address,
                "unit_number": unit_number,
            }
        )

        total_amount += invoice.amount

        if computed_status == InvoiceStatus.paid.value:
            paid_amount += invoice.amount
        elif computed_status == "overdue":
            overdue_amount += invoice.amount
        else:
            pending_amount += invoice.amount

    totals = {
        "count": len(rows),
        "total_amount": total_amount,
        "paid_amount": paid_amount,
        "pending_amount": pending_amount,
        "overdue_amount": overdue_amount,
    }

    return rows, totals


def build_tenant_space_report_rows(
    session: Session,
    tenant_ids: List[int],
    month_start: date,
    month_end: date,
) -> tuple[List[Dict[str, Any]], Dict[str, Any]]:
    leases = session.exec(
        select(Lease)
        .where(
            Lease.tenant_id.in_(tenant_ids),
            Lease.end_date >= month_start,
        )
        .order_by(Lease.end_date.desc())
    ).all()

    rows: List[Dict[str, Any]] = []

    total_area = decimal_zero()
    monthly_rent = decimal_zero()

    for lease in leases:
        if lease.start_date and lease.start_date > month_end:
            continue

        unit = session.get(Unit, lease.unit_id)

        if not unit:
            continue

        obj = session.get(Object, unit.object_id)

        if not obj:
            continue

        rows.append(
            {
                "object_address": obj.address,
                "unit_number": unit.number,
                "area": unit.area,
                "rent_monthly": lease.rent_monthly,
                "start_date": lease.start_date,
                "end_date": lease.end_date,
                "status": get_lease_status(lease.end_date),
            }
        )

        total_area += unit.area
        monthly_rent += lease.rent_monthly

    totals = {
        "count": len(rows),
        "total_area": total_area,
        "monthly_rent": monthly_rent,
    }

    return rows, totals


def get_tenant_report(
    session: Session,
    user_id: int,
    tab: str,
    period: str,
) -> TenantReportResponse:
    if tab not in ("invoices", "spaces"):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="tab must be either invoices or spaces",
        )

    month_start, month_end = parse_period(period)

    profile = get_tenant_profile(session, user_id)

    if not profile:
        return TenantReportResponse(
            tab=tab,
            period=period,
            matched=False,
            message=EMPTY_STATE_MESSAGE,
            rows=[],
            totals=empty_report_totals(tab),
        )

    tenant_ids = get_matched_tenant_ids(session, profile)

    if not tenant_ids:
        return TenantReportResponse(
            tab=tab,
            period=period,
            matched=False,
            message=EMPTY_STATE_MESSAGE,
            rows=[],
            totals=empty_report_totals(tab),
        )

    if tab == "invoices":
        rows, totals = build_tenant_invoice_report_rows(
            session,
            tenant_ids,
            period,
        )
    else:
        rows, totals = build_tenant_space_report_rows(
            session,
            tenant_ids,
            month_start,
            month_end,
        )

    return TenantReportResponse(
        tab=tab,
        period=period,
        matched=True,
        message=None,
        rows=rows,
        totals=totals,
    )


def get_tenant_subscription_response(
    session: Session,
    user_id: int,
) -> TenantSubscriptionResponse:
    subscription = get_or_create_subscription(
        session=session,
        user_id=user_id,
        role=Role.tenant.value,
    )

    return TenantSubscriptionResponse(
        plan=subscription.plan,
        status=subscription.status,
        started_at=subscription.started_at,
        expires_at=subscription.expires_at,
        can_export=subscription.plan == Plan.tenant_reports.value,
    )


def upgrade_tenant_subscription(
    session: Session,
    user_id: int,
) -> TenantSubscriptionResponse:
    subscription = get_or_create_subscription(
        session=session,
        user_id=user_id,
        role=Role.tenant.value,
    )

    subscription.plan = Plan.tenant_reports.value
    subscription.status = "active"

    session.add(subscription)
    session.commit()
    session.refresh(subscription)

    return get_tenant_subscription_response(session, user_id)


def normalize_tenant_export_fields(tab: str, fields: Optional[str]) -> List[str]:
    default_fields = INVOICE_REPORT_FIELDS if tab == "invoices" else SPACE_REPORT_FIELDS

    if not fields:
        return default_fields

    requested = [field.strip() for field in fields.split(",") if field.strip()]

    allowed = set(default_fields)

    result = [field for field in requested if field in allowed]

    return result or default_fields


def excel_cell_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)

    return value


def tenant_report_to_xlsx_bytes(
    report: TenantReportResponse,
    fields: Optional[str] = None,
) -> BytesIO:
    selected_fields = normalize_tenant_export_fields(report.tab, fields)

    wb = Workbook()
    ws = wb.active
    ws.title = "Tenant report"

    ws.append([FIELD_TITLES.get(field, field) for field in selected_fields])

    header_fill = PatternFill("solid", fgColor="1F4E78")

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in report.rows:
        ws.append([excel_cell_value(row.get(field)) for field in selected_fields])

    ws.append([])

    total_map = TOTAL_FIELD_MAP.get(report.tab, {})

    total_row = []

    for index, field in enumerate(selected_fields):
        if index == 0:
            total_row.append("Итого")
            continue

        total_key = total_map.get(field)

        if total_key and total_key in report.totals:
            total_row.append(excel_cell_value(report.totals.get(total_key)))
        else:
            total_row.append("")

    ws.append(total_row)

    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            value = str(cell.value) if cell.value is not None else ""
            max_length = max(max_length, len(value))

        ws.column_dimensions[column_letter].width = min(max_length + 2, 40)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output
